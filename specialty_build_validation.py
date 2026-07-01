"""Build the BLIND human-validation set for the SPECIALTY classifier.

Parallel to the Q7 validation (q7_build_validation.py + q7_build_coding_xlsx.py),
but for the 18-way clinical-specialty classifier. NOTHING is classified here —
we only SAMPLE existing committed outputs.

Reads (all from the repo root):
  output/specialty_classifications.csv   (pmid, specialty = MODEL label)
  output/classified_medical_Q1Q2.csv     (pmid, title, abstract, ...)
  q7_abstract_filter.abstract_present_pmids  (shared abstract-present helper)

Writes:
  output/analyses/specialty_validation_sample.xlsx  — BLIND coding sheet
      (columns: pmid | title | abstract | human_specialty; the model label is
       NOT present anywhere in this file). human_specialty is a data-validation
       DROPDOWN of the exact canonical taxonomy + "Unclear / cannot determine
       from abstract".
  output/analyses/specialty_validation_key.csv      — hidden key: pmid, model_specialty

Sampling (seed=42):
  1. restrict to abstract-present records (drop title-only, as in Q7)
  2. drop model labels in {FAILED, "", Unknown}
  3. stratify by MODEL-assigned specialty; draw up to ceil(120/n_specialties)
     per specialty, then top up at random to exactly n=120. If a specialty has
     fewer records than its quota, take all and redistribute the remainder.
"""
import os, csv, math
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from q7_abstract_filter import abstract_present_pmids

SEED   = 42
TARGET = 120
BASE   = os.path.dirname(os.path.abspath(__file__))
SPEC   = os.path.join(BASE, "output", "specialty_classifications.csv")
CORPUS = os.path.join(BASE, "output", "classified_medical_Q1Q2.csv")
OUT_XLSX = os.path.join(BASE, "output", "analyses", "specialty_validation_sample.xlsx")
OUT_KEY  = os.path.join(BASE, "output", "analyses", "specialty_validation_key.csv")

# ── canonical taxonomy — read VERBATIM from the committed specialty classifier ──
# analysis/classify_specialty.py (commit a30aaf7, branch analysis-session-2):
# these are the SHORT category names the classifier is told to emit and that it
# validates against (VALID_SPECIALTIES). Kept here as the single source of the
# dropdown menu. If this list ever drifts from the classifier, the classifier wins.
CANON_SPECIALTIES = [
    "Radiology / Diagnostic Imaging",
    "Pathology / Laboratory Medicine",
    "Cardiology",
    "Oncology",
    "Surgery",
    "Ophthalmology",
    "Dermatology",
    "Neurology / Neuroscience",
    "Mental Health / Psychiatry",
    "Internal Medicine / Primary Care",
    "Emergency Medicine / Critical Care",
    "Pediatrics",
    "Obstetrics / Gynecology",
    "Dentistry",
    "Medical Informatics / Digital Health",
    "Nursing",
    "Medical Education",
    "Multidisciplinary / Other",
]
UNCLEAR_OPTION = "Unclear / cannot determine from abstract"
DROPDOWN_OPTIONS = CANON_SPECIALTIES + [UNCLEAR_OPTION]

JUNK_LABELS = {"FAILED", "", "Unknown"}


# ── sampling ────────────────────────────────────────────────────────────────────
def allocate(avail, target, rng):
    """Return {specialty: n_to_draw} summing to min(target, total available).

    Water-filling with a randomized per-round order: every specialty is grown one
    at a time (so the allocation stays as even as possible and each specialty is
    capped near ceil(target/n)); a specialty that runs out of records simply drops
    out and its share is redistributed at random to the others. This realizes
    "up to ceil(target/n) per specialty, take all if short, redistribute + top up
    to exactly target at random".
    """
    specs = sorted(avail)
    alloc = {s: 0 for s in specs}
    total_avail = sum(avail.values())
    goal = min(target, total_avail)
    while sum(alloc.values()) < goal:
        order = specs[:]
        rng.shuffle(order)
        progressed = False
        for s in order:
            if sum(alloc.values()) >= goal:
                break
            if alloc[s] < avail[s]:
                alloc[s] += 1
                progressed = True
        if not progressed:
            break
    return alloc


def main():
    rng = np.random.RandomState(SEED)

    # 1. abstract-present restriction (shared helper, same as Q7)
    ap = abstract_present_pmids(BASE)

    # model labels
    spec_df = pd.read_csv(SPEC, dtype={"pmid": str})[["pmid", "specialty"]]
    spec_df["specialty"] = spec_df["specialty"].fillna("").astype(str).str.strip()

    # 2. eligible = abstract-present AND non-junk model label
    elig = spec_df[spec_df["pmid"].isin(ap) & ~spec_df["specialty"].isin(JUNK_LABELS)].copy()
    avail = elig["specialty"].value_counts().to_dict()
    n_spec = len(avail)
    quota = math.ceil(TARGET / n_spec)
    print(f"Abstract-present pmids: {len(ap):,}")
    print(f"Eligible records (abstract-present, non-junk model label): {len(elig):,}")
    print(f"Distinct model specialties: {n_spec}   nominal quota = ceil({TARGET}/{n_spec}) = {quota}")

    # 3. allocate then random-sample within each specialty
    alloc = allocate(avail, TARGET, rng)
    picks = []
    for s in sorted(alloc):
        k = alloc[s]
        if k == 0:
            continue
        pool = elig[elig["specialty"] == s]
        picks.append(pool.sample(k, random_state=rng))
    picked = pd.concat(picks).drop_duplicates("pmid")
    assert len(picked) == min(TARGET, len(elig)), f"got {len(picked)}, expected {TARGET}"

    # realized allocation report
    realized = picked["specialty"].value_counts()
    print(f"\nRealized per-specialty allocation (total = {len(picked)}):")
    for s in sorted(avail, key=lambda x: (-realized.get(x, 0), x)):
        print(f"  {realized.get(s, 0):>3d}  ({avail[s]:>5d} avail)  {s}")

    # 4. attach title + abstract from the corpus
    corpus = pd.read_csv(CORPUS, dtype={"pmid": str})[["pmid", "title", "abstract"]]
    key = picked[["pmid", "specialty"]].rename(columns={"specialty": "model_specialty"})
    sheet = picked[["pmid"]].merge(corpus, on="pmid", how="left")

    # shuffle so row order cannot leak the (hidden) model stratum
    sheet = sheet.sample(frac=1.0, random_state=SEED).reset_index(drop=True)
    sheet["human_specialty"] = ""
    sheet = sheet[["pmid", "title", "abstract", "human_specialty"]]

    # ── hidden key (kept separate from the blind sheet) ────────────────────────
    key.sort_values("pmid").to_csv(OUT_KEY, index=False)
    print(f"\nHidden key → {OUT_KEY}  ({len(key)} rows)")

    # ── blind coding workbook ──────────────────────────────────────────────────
    build_workbook(sheet)
    print(f"Blind coding sheet → {OUT_XLSX}  ({len(sheet)} rows)")
    print("Coder fills human_specialty from the dropdown; then run "
          "analysis/compute_specialty_kappa.py")


def build_workbook(sheet):
    FONT     = "Arial"
    HDR_FILL = PatternFill("solid", fgColor="305496")
    FILL_CODE = PatternFill("solid", fgColor="FFF2CC")   # yellow = fill me
    THIN   = Side(style="thin", color="BFBFBF")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    wb = Workbook()

    # ── Sheet 1: Read me first ─────────────────────────────────────────────────
    g = wb.active
    g.title = "Read me first"
    g.sheet_view.showGridLines = False
    g.column_dimensions["A"].width = 34
    g.column_dimensions["B"].width = 96

    def line(row, a, b="", *, bold=False, size=11, color="000000", wrap=True):
        ca, cb = g.cell(row=row, column=1, value=a), g.cell(row=row, column=2, value=b)
        for c in (ca, cb):
            c.font = Font(name=FONT, bold=bold, size=size, color=color)
            c.alignment = Alignment(vertical="top", wrap_text=wrap)
        return row + 1

    r = 1
    r = line(r, "Specialty validation — blind human coding", "",
             bold=True, size=15, color="305496")
    r = line(r, "What to do",
             "For each paper on the Coding tab, read the title + abstract and pick ONE "
             "clinical specialty in the yellow human_specialty drop-down. These are your "
             "independent judgments; the model's predicted specialty is deliberately NOT "
             "shown. Assign by the clinical specialty being applied to or studied — not by "
             "the AI methodology used.")
    r += 1
    r = line(r, "The 18 categories", "(pick the single best fit)", bold=True, size=12, color="C00000")
    rules = {
        "Radiology / Diagnostic Imaging": "imaging: X-ray, CT, MRI, ultrasound, PET; image classification / segmentation.",
        "Pathology / Laboratory Medicine": "histopathology, lab tests, digital pathology, biomarkers.",
        "Cardiology": "heart & cardiovascular (non-surgical).",
        "Oncology": "cancer — takes priority over organ system (e.g. breast-cancer detection → Oncology).",
        "Surgery": "general surgery AND all surgical subspecialties (orthopedic, ENT, urology, neurosurgery, plastic, ...).",
        "Ophthalmology": "eye / vision, retinal imaging.",
        "Dermatology": "skin, dermoscopy, skin-lesion classification.",
        "Neurology / Neuroscience": "non-surgical neuro; brain/nervous-system disease, neuroimaging read as neurology.",
        "Mental Health / Psychiatry": "psychiatry, mental health, psychology.",
        "Internal Medicine / Primary Care": "family/primary care + all non-surgical internal subspecialties (GI, endocrine, rheum, nephro, pulm, ID).",
        "Emergency Medicine / Critical Care": "ED, ICU, critical / intensive care, triage.",
        "Pediatrics": "children — takes priority over organ system (pediatric cardiology → Pediatrics).",
        "Obstetrics / Gynecology": "pregnancy, obstetrics, gynecology.",
        "Dentistry": "dental, oral / maxillofacial, orthodontics.",
        "Medical Informatics / Digital Health": "papers ABOUT AI itself — frameworks, taxonomies, benchmarks, general 'AI in medicine' with no single clinical anchor.",
        "Nursing": "nursing practice / education.",
        "Medical Education": "training of clinicians / students; AI literacy; exams (e.g. 'students' AI literacy' → Medical Education).",
        "Multidisciplinary / Other": "spans 2+ specialties with no dominant focus; bibliometric reviews; veterinary; rehab / anesthesiology / public health when not clearly elsewhere.",
    }
    for name in CANON_SPECIALTIES:
        r = line(r, name, rules[name])
    r = line(r, UNCLEAR_OPTION,
             "Use ONLY if the title + abstract genuinely do not let you place the paper in any "
             "of the 18 categories. Reported separately as a sensitivity analysis.")
    r += 1
    r = line(r, "Priority rules", "", bold=True, size=12, color="C00000")
    r = line(r, "Oncology & Pediatrics win", "Cancer → Oncology; children → Pediatrics, over the organ-system category.")
    r = line(r, "Specialty, not method", "'ChatGPT for cardiology questions' → Cardiology, not Medical Informatics.")
    r = line(r, "Informatics is for AI-about-AI", "Reserve Medical Informatics for papers whose subject is the AI/method itself.")

    for row in range(1, r):
        g.row_dimensions[row].height = 30

    # ── Sheet 2: Coding (exactly pmid | title | abstract | human_specialty) ─────
    s = wb.create_sheet("Coding")
    s.sheet_view.showGridLines = False
    headers = ["pmid", "title", "abstract", "human_specialty"]
    widths  = [12, 55, 110, 34]
    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        c = s.cell(row=1, column=j, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        s.column_dimensions[get_column_letter(j)].width = w

    for i, rec in enumerate(sheet.itertuples(index=False), start=2):
        vals = [rec.pmid, rec.title, rec.abstract, ""]
        for j, v in enumerate(vals, start=1):
            c = s.cell(row=i, column=j, value=("" if pd.isna(v) else v))
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top",
                                    wrap_text=(j in (2, 3)),
                                    horizontal="center" if j == 1 else "left")
            c.border = BORDER
            if j == 4:
                c.fill = FILL_CODE
        s.row_dimensions[i].height = 150

    # ── hidden helper sheet holding the dropdown menu ──────────────────────────
    # (an inline list would exceed Excel's ~255-char data-validation limit, so we
    #  point the validation at a real cell range instead.)
    lst = wb.create_sheet("lists")
    for i, opt in enumerate(DROPDOWN_OPTIONS, start=1):
        lst.cell(row=i, column=1, value=opt)
    lst.sheet_state = "hidden"
    n_opt = len(DROPDOWN_OPTIONS)

    last = len(sheet) + 1
    dv = DataValidation(type="list", formula1=f"lists!$A$1:$A${n_opt}",
                        allow_blank=True, showDropDown=False)
    dv.error = "Pick one specialty from the dropdown."
    dv.errorTitle = "Invalid specialty"
    dv.prompt = "Choose ONE specialty (or 'Unclear ...')."
    dv.promptTitle = "human_specialty"
    s.add_data_validation(dv)
    dv.add(f"D2:D{last}")

    s.freeze_panes = "A2"            # freeze header row
    s.auto_filter.ref = f"A1:D{last}"

    wb.save(OUT_XLSX)


if __name__ == "__main__":
    main()
