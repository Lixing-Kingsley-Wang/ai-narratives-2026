#!/usr/bin/env python3
"""
Prophecy Panel — tier -> verdict Sankey (protocol A8), exported as publication-grade SVG.

Left  = 3 tiers (node height proportional to n; falsifiability gradient T1 darkest -> T3 lightest).
Right = 4 verdicts (node height proportional to n).
Every ribbon width proportional to count and LABELLED with its count. Small n not masked.
Reads verdicts from phase2_prophecy_verdicts_worksheet.xlsx (source of truth).
"""
import os, sys
from collections import Counter, defaultdict
from openpyxl import load_workbook

OUT_DIR = "/Users/kingslywang/repos/ai-narratives-2026/output/analyses/prophecy"
XLSX = os.path.join(OUT_DIR, "phase2_prophecy_verdicts_worksheet.xlsx")
SVG  = os.path.join(OUT_DIR, "prophecy_sankey.svg")

TIERS = ["1", "2", "3"]
TIER_LABEL = {"1": "Tier 1\nSpecific superiority /\nreplacement",
              "2": "Tier 2\nCapability\ndeployment",
              "3": "Tier 3\nDiffuse\ntransformation"}
VERDICTS = ["borne_out", "partially", "not_borne_out", "too_early_unfalsifiable"]
VERDICT_LABEL = {"borne_out": "Borne out", "partially": "Partially borne out",
                 "not_borne_out": "Not borne out", "too_early_unfalsifiable": "Too early /\nunfalsifiable"}
TIER_FILL = {"1": "#334155", "2": "#64748B", "3": "#94A3B8"}          # dark -> light gradient
VERDICT_FILL = {"borne_out": "#2E7D32", "partially": "#CFA85A",
                "not_borne_out": "#C23B22", "too_early_unfalsifiable": "#5F6C8A"}
# ribbons coloured by DESTINATION verdict. Three hero flows emphasised; all others muted.
HERO = {("1", "not_borne_out"): 0.80, ("2", "too_early_unfalsifiable"): 0.75, ("3", "partially"): 0.40}
DEFAULT_ALPHA = 0.25

# ── geometry ─────────────────────────────────────────────────────────────────────
SCALE = 5.6          # px per claim
GAP   = 26           # px between stacked nodes
NW    = 24           # node width
X0    = 250          # left-node left edge
X1    = 700          # right-node left edge
LEFT_RIGHT_EDGE = X0 + NW
W, H  = 1180, 700
CONTENT_TOP = 104    # ribbons/nodes live in [CONTENT_TOP, CONTENT_BOT]; caption sits below
CONTENT_BOT = 648


def esc(s): return s.replace("&", "&amp;").replace("<", "&lt;")


def stack(order, counts):
    total = sum(counts[k] for k in order)
    span = total * SCALE + (len([k for k in order if counts[k]]) - 1) * GAP
    y = CONTENT_TOP + (CONTENT_BOT - CONTENT_TOP - span) / 2
    pos = {}
    for k in order:
        if counts[k]:
            h = counts[k] * SCALE
            pos[k] = (y, y + h); y += h + GAP
    return pos


def curve(x0, y0, x1, y1):
    xm = (x0 + x1) / 2
    return f"M {x0:.1f} {y0:.1f} C {xm:.1f} {y0:.1f} {xm:.1f} {y1:.1f} {x1:.1f} {y1:.1f}"


def ribbon(x0, at, ab, x1, bt, bb, fill, alpha):
    xm = (x0 + x1) / 2
    d = (f"M {x0:.1f} {at:.1f} C {xm:.1f} {at:.1f} {xm:.1f} {bt:.1f} {x1:.1f} {bt:.1f} "
         f"L {x1:.1f} {bb:.1f} C {xm:.1f} {bb:.1f} {xm:.1f} {ab:.1f} {x0:.1f} {ab:.1f} Z")
    return f'<path d="{d}" fill="{fill}" fill-opacity="{alpha}"/>'


def main():
    s = load_workbook(XLSX, data_only=True)["adjudicate"]
    H_ = {s.cell(1, j).value: j for j in range(1, s.max_column + 1)}
    def c(r, n):
        v = s.cell(r, H_[n]).value; return "" if v is None else str(v).strip()
    rows = [r for r in range(2, s.max_row + 1) if c(r, "claim_id")]
    M = Counter((c(r, "tier"), c(r, "verdict")) for r in rows)
    tier_n = {t: sum(M[(t, v)] for v in VERDICTS) for t in TIERS}
    verd_n = {v: sum(M[(t, v)] for t in TIERS) for v in VERDICTS}
    N = len(rows)

    lpos = stack(TIERS, tier_n)
    rpos = stack(VERDICTS, verd_n)

    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}" font-family="Helvetica,Arial,sans-serif">']
    svg.append(f'<rect width="{W}" height="{H}" fill="#ffffff"/>')
    svg.append(f'<text x="{W/2:.0f}" y="40" text-anchor="middle" font-size="20" font-weight="700" fill="#1a1a1a">Prophecy Panel: claim tier &#8594; adjudicated verdict (N = {N})</text>')
    svg.append(f'<text x="{W/2:.0f}" y="62" text-anchor="middle" font-size="12" fill="#666">Node height &#8733; number of claims; every ribbon width &#8733; its count. 2021&#8211;2023 Advocacy predictions vs the 2024&#8211;2026 record.</text>')

    # ribbons (draw first, under nodes). Left sub-bands ordered by verdict; right by tier.
    loff = {t: lpos[t][0] for t in TIERS if t in lpos}
    roff = {v: rpos[v][0] for v in VERDICTS if v in rpos}
    labels = []
    for t in TIERS:
        for v in VERDICTS:
            n = M[(t, v)]
            if not n:
                continue
            hgt = n * SCALE
            at, ab = loff[t], loff[t] + hgt; loff[t] += hgt
            bt, bb = roff[v], roff[v] + hgt; roff[v] += hgt
            svg.append(ribbon(LEFT_RIGHT_EDGE, at, ab, X1, bt, bb, VERDICT_FILL[v], HERO.get((t, v), DEFAULT_ALPHA)))
            # label sits near the TARGET node (right), at the height where the ribbon lands
            labels.append((X1 - 30, (bt + bb) / 2, n))

    # ribbon count labels (on top, with halo); small flows smaller
    for lx, ly, n in labels:
        fs = 10 if n < 5 else 12
        svg.append(f'<text x="{lx:.1f}" y="{ly:.1f}" text-anchor="middle" dominant-baseline="middle" '
                   f'font-size="{fs}" font-weight="700" fill="#111" stroke="#ffffff" stroke-width="2.5" paint-order="stroke">{n}</text>')

    # tier nodes + labels
    for t in TIERS:
        if t not in lpos: continue
        y0, y1 = lpos[t]
        svg.append(f'<rect x="{X0}" y="{y0:.1f}" width="{NW}" height="{y1-y0:.1f}" rx="3" fill="{TIER_FILL[t]}"/>')
        cy = (y0 + y1) / 2
        lines = TIER_LABEL[t].split("\n")
        ly = cy - (len(lines) - 1) * 7
        for k, ln in enumerate(lines):
            fw = "700" if k == 0 else "400"; fs = 13 if k == 0 else 11
            svg.append(f'<text x="{X0-12}" y="{ly + k*14:.1f}" text-anchor="end" font-size="{fs}" font-weight="{fw}" fill="#1a1a1a">{esc(ln)}</text>')
        svg.append(f'<text x="{X0-12}" y="{ly + len(lines)*14:.1f}" text-anchor="end" font-size="11" font-weight="700" fill="{TIER_FILL[t]}">n = {tier_n[t]}</text>')

    # verdict nodes + labels
    for v in VERDICTS:
        if v not in rpos: continue
        y0, y1 = rpos[v]
        svg.append(f'<rect x="{X1}" y="{y0:.1f}" width="{NW}" height="{y1-y0:.1f}" rx="3" fill="{VERDICT_FILL[v]}"/>')
        cy = (y0 + y1) / 2
        lines = VERDICT_LABEL[v].split("\n")
        ly = cy - (len(lines) - 1) * 7
        for k, ln in enumerate(lines):
            svg.append(f'<text x="{X1+NW+12}" y="{ly + k*14:.1f}" text-anchor="start" font-size="13" font-weight="600" fill="#1a1a1a">{esc(ln)}</text>')
        svg.append(f'<text x="{X1+NW+12}" y="{ly + len(lines)*14:.1f}" text-anchor="start" font-size="11" font-weight="700" fill="{VERDICT_FILL[v]}">n = {verd_n[v]}</text>')

    svg.append(f'<text x="{X0}" y="{H-24}" font-size="11" fill="#888">Falsifiability gradient: Tier 1 (darkest) &#8594; Tier 3 (lightest).</text>')
    svg.append("</svg>")
    open(SVG, "w", encoding="utf-8").write("\n".join(svg))
    print(f"wrote {SVG}")
    print("tier_n:", tier_n, "verd_n:", verd_n, "N:", N)
    print("matrix:")
    for t in TIERS:
        print(f"  T{t}: " + "  ".join(f"{v}={M[(t,v)]}" for v in VERDICTS if M[(t,v)]))


if __name__ == "__main__":
    main()
