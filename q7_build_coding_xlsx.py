"""Build the human coding workbook (drop-downs) from q7_validation_blind.csv.

Sheet 1 "Read me first" — axis definitions, decision rules, calibration examples.
Sheet 2 "Coding"       — 120 papers; human_failure_mode / human_model_type are
                         drop-down (data-validation) columns, blank to fill.
Model labels are NOT included (blind coding).
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
BLIND = os.path.join(BASE, "output", "analyses", "q7_validation_blind.csv")
OUT   = os.path.join(BASE, "output", "analyses", "q7_validation_coding.xlsx")

FM = ["confabulation", "misclassification", "both", "none_or_unclear"]
MT = ["generative", "discriminative", "both", "unclear"]

FONT      = "Arial"
HDR_FILL  = PatternFill("solid", fgColor="305496")
FILL_CODE = PatternFill("solid", fgColor="FFF2CC")   # yellow = fill me
THIN      = Side(style="thin", color="BFBFBF")
BORDER    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

df = pd.read_csv(BLIND, dtype={"pmid": str})

wb = Workbook()

# ── Sheet 1: Read me first ─────────────────────────────────────────────────────
g = wb.active
g.title = "Read me first"
g.sheet_view.showGridLines = False
g.column_dimensions["A"].width = 24
g.column_dimensions["B"].width = 95

def line(row, a, b="", *, bold=False, size=11, color="000000", fill=None, wrap=True):
    ca, cb = g.cell(row=row, column=1, value=a), g.cell(row=row, column=2, value=b)
    for c in (ca, cb):
        c.font = Font(name=FONT, bold=bold, size=size, color=color)
        c.alignment = Alignment(vertical="top", wrap_text=wrap)
        if fill: c.fill = fill
    return row + 1

r = 1
r = line(r, "Q7 validation — blind human coding", "", bold=True, size=15, color="305496")
r = line(r, "What to do",
         "For each paper (Coding tab), read the title + abstract and choose ONE value in each of the two "
         "yellow drop-down columns. Code from the MECHANISM described — not from whether a brand name "
         "(GPT / ChatGPT / etc.) appears. These are your independent judgments; the model's labels are "
         "deliberately NOT shown.", bold=False)
r += 1
r = line(r, "Axis A — failure_mode", "(what kind of failure does the paper describe?)", bold=True, size=12, color="C00000")
r = line(r, "confabulation", "AI GENERATES fabricated / unfaithful content: invented citations or references, "
         "made-up facts, plausible-but-false text, hallucinated entities, unfaithful summaries.")
r = line(r, "misclassification", "AI produces a WRONG label / score / prediction / detection: false positives or "
         "negatives, misdiagnosis, miscalibration, poor generalization, dataset shift, biased predictions, "
         "wrong answers on a test.")
r = line(r, "both", "Both mechanisms are explicitly described in the same paper.")
r = line(r, "none_or_unclear", "No specific mechanism is described, or only a general 'AI risks / needs oversight' framing.")
r += 1
r = line(r, "Axis B — model_type", "(what kind of AI is being evaluated?)", bold=True, size=12, color="C00000")
r = line(r, "generative", "LLM, chatbot, or text / image / code generator (e.g. ChatGPT, GPT-4, a diffusion image model).")
r = line(r, "discriminative", "Predictive / classification / detection / segmentation / risk-scoring model "
         "(e.g. a CNN reading mammograms, a survival-risk model).")
r = line(r, "both", "Both families are explicitly evaluated in the same paper.")
r = line(r, "unclear", "Cannot tell from title + abstract.")
r += 1
r = line(r, "Key decision rules", "", bold=True, size=12, color="C00000")
r = line(r, "Mechanism, not brand", "An LLM tested for ACCURACY is misclassification, NOT confabulation. "
         "Fabrication only = confabulation.")
r = line(r, "High accuracy ≠ fabrication", "'ChatGPT scored 85% on an exam' → misclassification (it's an accuracy claim).")
r = line(r, "'Hallucination'", "When a paper uses 'hallucination' to mean the model MADE SOMETHING UP "
         "(invented refs/facts), that is confabulation. If it just means 'wrong', that is misclassification.")
r = line(r, "No mechanism", "General 'AI is risky, needs oversight' with no described failure → "
         "none_or_unclear + unclear.")
r = line(r, "When genuinely unsure", "Pick none_or_unclear / unclear and write a word in Notes. Do not guess a specific mechanism.")
r += 1
r = line(r, "Calibration examples", "", bold=True, size=12, color="C00000")
ex = [
    ("ChatGPT scored 85% on USMLE-style questions", "misclassification + generative"),
    ("Deep-learning classifier produced false-negative mammograms", "misclassification + discriminative"),
    ("LLM discharge summary fabricated a drug and an invented citation", "confabulation + generative"),
    ("Measured GPT-4 hallucinated references AND its diagnostic accuracy", "both + generative"),
    ("AI errors threaten safety, need oversight (no mechanism given)", "none_or_unclear + unclear"),
]
for paper, label in ex:
    r = line(r, f"“{paper}”", f"→  {label}")

for row in range(1, r):
    g.row_dimensions[row].height = 30

# ── Sheet 2: Coding ────────────────────────────────────────────────────────────
s = wb.create_sheet("Coding")
s.sheet_view.showGridLines = False
headers = ["row", "pmid", "pub_year", "title", "abstract",
           "human_failure_mode", "human_model_type", "notes"]
widths  = [5, 11, 9, 42, 90, 20, 20, 26]
for j, (h, w) in enumerate(zip(headers, widths), start=1):
    c = s.cell(row=1, column=j, value=h)
    c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
    c.fill = HDR_FILL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    s.column_dimensions[get_column_letter(j)].width = w

for i, rec in enumerate(df.itertuples(index=False), start=2):
    vals = [i - 1, rec.pmid, rec.pub_year, rec.title, rec.abstract, "", "", ""]
    for j, v in enumerate(vals, start=1):
        c = s.cell(row=i, column=j, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(vertical="top",
                                wrap_text=(j in (4, 5, 8)),
                                horizontal="center" if j in (1, 3, 6, 7) else "left")
        c.border = BORDER
        if j in (6, 7):
            c.fill = FILL_CODE
    s.row_dimensions[i].height = 150

# drop-downs (whole column range)
last = len(df) + 1
dv_fm = DataValidation(type="list", formula1='"%s"' % ",".join(FM),
                       allow_blank=True, showDropDown=False)
dv_fm.error = "Pick one: " + ", ".join(FM); dv_fm.errorTitle = "Invalid failure_mode"
dv_fm.prompt = "Choose ONE failure_mode"; dv_fm.promptTitle = "failure_mode"
dv_mt = DataValidation(type="list", formula1='"%s"' % ",".join(MT),
                       allow_blank=True, showDropDown=False)
dv_mt.error = "Pick one: " + ", ".join(MT); dv_mt.errorTitle = "Invalid model_type"
dv_mt.prompt = "Choose ONE model_type"; dv_mt.promptTitle = "model_type"
s.add_data_validation(dv_fm); s.add_data_validation(dv_mt)
dv_fm.add(f"F2:F{last}")
dv_mt.add(f"G2:G{last}")

s.freeze_panes = "A2"
s.auto_filter.ref = f"A1:H{last}"

wb.save(OUT)
print(f"Saved coding workbook → {OUT}")
print(f"  Coding tab: {len(df)} papers; drop-downs on F (failure_mode) and G (model_type).")
print(f"  failure_mode options: {FM}")
print(f"  model_type   options: {MT}")
